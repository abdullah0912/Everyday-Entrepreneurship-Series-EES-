import os
import tkinter as tk
from tkinter import messagebox
import openpyxl

class LabelPrintApp(tk.Tk):
    def __init__(self):
        super().__init__()
        # Pencere Başlığı (Sade)
        self.title("Label Print Automation")
        
        # Ekranın tam ortasında aç
        self.merkeze_hizala(540, 480)
        
        self.configure(bg="#f8f9fa")
        self.resizable(False, False)

        # Excel dosya yolu
        self.excel_yolu = os.path.join(os.path.dirname(__file__), "data", "etiket_sablonu.xlsx")
        
        self.secilen_indirim = None
        self.urun_listesi = []

        # Arayüz Konteyneri
        self.govde = tk.Frame(self, bg="#f8f9fa")
        self.govde.pack(fill="both", expand=True)

        # Başlangıçta Excel verilerini hafızaya al
        self.verileri_yukle()
        self.sayfa_oran_secimi()

    def merkeze_hizala(self, genislik, yukseklik):
        """Pencereyi ekran çözünürlüğüne göre tam merkeze yerleştirir."""
        self.update_idletasks()
        ekran_genislik = self.winfo_screenwidth()
        ekran_yukseklik = self.winfo_screenheight()

        x = int((ekran_genislik / 2) - (genislik / 2))
        y = int((ekran_yukseklik / 2) - (yukseklik / 2))

        self.geometry(f"{genislik}x{yukseklik}+{x}+{y}")

    def verileri_yukle(self):
        """Excel'deki 'SIM STİL YÖNETİMİ RAPORU' sayfasını listeye çeker."""
        if not os.path.exists(self.excel_yolu):
            print("HATA: Excel dosyası bulunamadı!")
            return

        try:
            wb = openpyxl.load_workbook(self.excel_yolu, data_only=True)
            if "SIM STİL YÖNETİMİ RAPORU" not in wb.sheetnames:
                print("Sayfa bulunamadı!")
                return
                
            ws = wb["SIM STİL YÖNETİMİ RAPORU"]
            self.urun_listesi.clear()

            for row in range(2, ws.max_row + 1):
                kod = str(ws.cell(row=row, column=1).value or "").strip() # A sütunu
                if kod:
                    ad = str(ws.cell(row=row, column=2).value or "").strip() # B sütunu
                    fiyat_ham = ws.cell(row=row, column=8).value # H sütunu
                    self.urun_listesi.append({
                        "kod": kod,
                        "ad": ad,
                        "fiyat": fiyat_ham
                    })
            print(f"Toplam {len(self.urun_listesi)} ürün başarıyla yüklendi!")
        except Exception as e:
            print(f"Excel okunurken hata oluştu: {e}")

    def temizle(self):
        for w in self.govde.winfo_children():
            w.destroy()

    def sayfa_oran_secimi(self):
        self.temizle()

        # Tam Ortalanmış Başlık Alanı
        baslik_frame = tk.Frame(self.govde, bg="#f8f9fa")
        baslik_frame.pack(pady=(45, 15))

        tk.Label(
            baslik_frame, 
            text="🏷️ Label Print Automation", 
            font=("Segoe UI", 18, "bold"), 
            bg="#f8f9fa", fg="#111"
        ).pack()

        tk.Label(
            self.govde, 
            text="Uygulanacak kampanya oranını seçiniz:", 
            font=("Segoe UI", 11), 
            bg="#f8f9fa", fg="#555"
        ).pack(pady=(0, 25))

        btn_kapsayici = tk.Frame(self.govde, bg="#f8f9fa")
        btn_kapsayici.pack()

        oranlar = [
            ("%20 İndirim", 20, "#0d6efd"),
            ("%30 İndirim", 30, "#fd7e14"),
            ("%50 İndirim", 50, "#dc3545")
        ]

        for text, oran, renk in oranlar:
            btn = tk.Button(
                btn_kapsayici, text=text, font=("Segoe UI", 11, "bold"),
                bg=renk, fg="white", width=14, height=2, relief="flat", cursor="hand2",
                command=lambda o=oran: self.oran_kaydet(o)
            )
            btn.pack(pady=6)

        # Bilgi Notu
        tk.Label(
            self.govde,
            text=f"📦 Sistemde kayıtlı ürün sayısı: {len(self.urun_listesi)}",
            font=("Segoe UI", 9, "italic"),
            bg="#f8f9fa", fg="#6c757d"
        ).pack(side="bottom", pady=15)

    def oran_kaydet(self, oran):
        self.secilen_indirim = oran
        self.sayfa_barkod_giris()

    def sayfa_barkod_giris(self):
        self.temizle()

        # Ortalanmış Üst Başlık
        baslik_frame = tk.Frame(self.govde, bg="#f8f9fa")
        baslik_frame.pack(pady=(20, 5))

        tk.Label(
            baslik_frame, 
            text="🏷️ Label Print Automation", 
            font=("Segoe UI", 15, "bold"), 
            bg="#f8f9fa", fg="#212529"
        ).pack()

        tk.Label(
            self.govde, 
            text=f"Aktif Kampanya: %{self.secilen_indirim} İndirim", 
            font=("Segoe UI", 12, "bold"), 
            bg="#f8f9fa", fg="#dc3545"
        ).pack(pady=(0, 10))

        tk.Label(
            self.govde, 
            text="Model Kodu / Barkod Giriniz (veya Okutunuz):", 
            font=("Segoe UI", 10), 
            bg="#f8f9fa"
        ).pack(pady=5)

        self.txt_barkod = tk.Entry(self.govde, font=("Segoe UI", 15), justify="center", width=22)
        self.txt_barkod.pack(pady=6, ipady=4)
        self.txt_barkod.focus_set()

        self.txt_barkod.bind("<Return>", self.urunu_ara_ve_goster)

        # Bilgi Kartı Alanı
        self.kart_frame = tk.Frame(self.govde, bg="#ffffff", bd=1, relief="solid", padx=15, pady=10)
        self.kart_frame.pack(pady=12, fill="x", padx=40)

        self.lbl_kart_model = tk.Label(self.kart_frame, text="Model: Bekleniyor...", font=("Segoe UI", 10, "bold"), bg="#ffffff", anchor="w")
        self.lbl_kart_model.pack(fill="x")

        self.lbl_kart_fiyat = tk.Label(self.kart_frame, text="Fiyat: -", font=("Segoe UI", 10), bg="#ffffff", anchor="w")
        self.lbl_kart_fiyat.pack(fill="x")

        self.lbl_kart_indirimli = tk.Label(self.kart_frame, text="İndirimli Tutar: -", font=("Segoe UI", 11, "bold"), fg="#198754", bg="#ffffff", anchor="w")
        self.lbl_kart_indirimli.pack(fill="x")

        btn_geri = tk.Button(
            self.govde, text="← Oran Değiştir", font=("Segoe UI", 9), relief="groove",
            command=self.sayfa_oran_secimi
        )
        btn_geri.pack(side="bottom", pady=15)

    def urunu_ara_ve_goster(self, event=None):
        kod = self.txt_barkod.get().strip()
        if not kod:
            return

        bulunan = None
        for urun in self.urun_listesi:
            if kod.lower() in urun["kod"].lower():
                bulunan = urun
                break

        if bulunan:
            try:
                ham = str(bulunan['fiyat']).replace("₺", "").replace(".", "").replace(",", ".").strip()
                eski_fiyat = float(ham)
                yeni_fiyat = eski_fiyat * (1 - (self.secilen_indirim / 100))

                self.lbl_kart_model.config(text=f"Model: {bulunan['kod']} - {bulunan['ad']}", fg="#111")
                self.lbl_kart_fiyat.config(text=f"Eski Fiyat: ₺{eski_fiyat:,.2f}")
                self.lbl_kart_indirimli.config(text=f"%{self.secilen_indirim} İndirimli: ₺{yeni_fiyat:,.2f}", fg="#198754")
            except Exception:
                self.lbl_kart_model.config(text=f"Model: {bulunan['kod']}", fg="#111")
                self.lbl_kart_fiyat.config(text=f"Orijinal Değer: {bulunan['fiyat']}")
                self.lbl_kart_indirimli.config(text="Hesaplanamadı", fg="#dc3545")
        else:
            self.lbl_kart_model.config(text="⚠️ Ürün Bulunamadı!", fg="#dc3545")
            self.lbl_kart_fiyat.config(text="-")
            self.lbl_kart_indirimli.config(text="-")

        self.txt_barkod.delete(0, tk.END)

if __name__ == "__main__":
    app = LabelPrintApp()
    app.mainloop()